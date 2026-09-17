"""Run a small real-data slice to estimate full-run cost and audit behavior."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from pbs_model import PBSDispatcher, detailed_scores, read_vehicles


def main() -> None:
    run_root = Path(__file__).resolve().parents[1]
    source = run_root / "inputs-raw" / "attachment1.xlsx"
    vehicles = read_vehicles(source)[:20]
    dispatcher = PBSDispatcher(vehicles, scenario="Q1", strategy="greedy")
    result = dispatcher.run()
    print("feasible", result.feasible)
    print("violations", result.violations)
    print("output", result.output_order)
    print("completion", result.completion_time, "returns", result.return_use_count)
    print("scores", detailed_scores(result, dispatcher.vehicle_by_id))
    print("events", len(result.event_log))
    print("timeline_cells", sum(len(value) for value in result.timeline.values()))
    workbook = Workbook()
    worksheet = workbook.active
    for row_index, vehicle_id in enumerate(result.output_order, 1):
        worksheet.cell(row=row_index, column=1, value=vehicle_id)
        worksheet.cell(row=row_index, column=2, value=result.completion_time)
    workbook.save(run_root / "work" / "small-slice.xlsx")


if __name__ == "__main__":
    main()
