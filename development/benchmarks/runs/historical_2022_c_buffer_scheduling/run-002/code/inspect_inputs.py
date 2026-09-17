"""Print compact inspection views for the audited 2022 C inputs."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def print_vehicle_file(path: Path) -> None:
    worksheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    header = rows[0]
    body = rows[1:]
    print(f"\n### {path.name}")
    print("header:", header)
    print("data rows:", len(body))
    print("first 8:", body[:8])
    print("last 8:", body[-8:])
    print("missing by column:", [sum(row[i] is None for row in body) for i in range(4)])
    for index, name in enumerate(header):
        counts = Counter(row[index] for row in body)
        print(f"{index}: {name}: {counts.most_common(20)}")
    ids = [row[0] for row in body]
    print("ids unique:", len(ids) == len(set(ids)), "min/max:", min(ids), max(ids))


def print_table(path: Path, limit: int = 100) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    print(f"\n### {path.name}")
    for worksheet in workbook.worksheets:
        rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        print("sheet:", worksheet.title, "dimension:", worksheet.calculate_dimension())
        for index, row in enumerate(rows[:limit], 1):
            print(index, row)
        if len(rows) > limit:
            print("... last 5 ...")
            for index, row in enumerate(rows[-5:], len(rows) - 4):
                print(index, row)


def main() -> None:
    root = Path(__file__).resolve().parents[1] / "inputs-raw"
    print_vehicle_file(root / "attachment1.xlsx")
    print_vehicle_file(root / "attachment2.xlsx")
    print_table(root / "attachment3.xlsx")
    print_table(root / "attachment4.xlsx", limit=5)


if __name__ == "__main__":
    main()
