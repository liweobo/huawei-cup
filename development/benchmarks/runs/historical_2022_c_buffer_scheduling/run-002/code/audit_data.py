"""Read-only structural audit for the 2022 C PBS scheduling attachments."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    return str(value)


def range_summary(values: list[Any]) -> dict[str, Any]:
    numeric = [value for value in values if isinstance(value, (int, float))]
    return {
        "count": len(values),
        "numeric_count": len(numeric),
        "min": min(numeric) if numeric else None,
        "max": max(numeric) if numeric else None,
        "unique_count": len(set(values)),
    }


def audit_workbook(path: Path, max_examples: int) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    report: dict[str, Any] = {
        "filename": path.name,
        "sha256": sha256(path),
        "bytes": path.stat().st_size,
        "sheets": [],
    }
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        nonempty_rows = [
            row for row in rows if any(value is not None for value in row)
        ]
        header = list(nonempty_rows[0]) if nonempty_rows else []
        data_rows = nonempty_rows[1:] if nonempty_rows else []
        width = max((len(row) for row in nonempty_rows), default=0)
        normalized = [
            tuple(row) + (None,) * (width - len(row)) for row in nonempty_rows
        ]
        columns = []
        for index in range(width):
            values = [row[index] for row in normalized[1:]]
            present = [value for value in values if value is not None]
            counts = Counter(str(value) for value in present)
            columns.append(
                {
                    "index": index + 1,
                    "header": header[index] if index < len(header) else None,
                    "nonempty": len(present),
                    "missing": len(values) - len(present),
                    "range": range_summary(present),
                    "top_values": counts.most_common(12),
                    "examples": json_safe(present[:max_examples]),
                }
            )
        duplicate_full_rows = len(data_rows) - len(set(normalized[1:]))
        report["sheets"].append(
            {
                "name": worksheet.title,
                "dimensions": worksheet.calculate_dimension(),
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "nonempty_row_count": len(nonempty_rows),
                "data_row_count": len(data_rows),
                "header": json_safe(header),
                "duplicate_full_rows": duplicate_full_rows,
                "columns": columns,
                "first_rows": json_safe(nonempty_rows[: min(12, len(nonempty_rows))]),
            }
        )
    workbook.close()
    return report


def build_consistency(audits: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {
        "attachment1_vs_attachment2": {},
        "attachment3_region_code_checks": {},
        "attachment4_template_checks": {},
    }
    by_name = {audit["filename"]: audit for audit in audits}
    data_sheets: dict[str, list[list[Any]]] = {}
    for name in ("attachment1.xlsx", "attachment2.xlsx"):
        sheet = by_name[name]["sheets"][0]
        rows = sheet["first_rows"]
        # Full rows are needed for exact comparisons; re-read below.
        data_sheets[name] = rows

    # Load complete cell values for the two vehicle lists.
    root = Path(__file__).resolve().parents[1] / "inputs-raw"
    for name in ("attachment1.xlsx", "attachment2.xlsx"):
        workbook = load_workbook(root / name, read_only=True, data_only=True)
        worksheet = workbook.worksheets[0]
        data_sheets[name] = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()

    a1 = data_sheets["attachment1.xlsx"]
    a2 = data_sheets["attachment2.xlsx"]
    a1_header = a1[0] if a1 else []
    a2_header = a2[0] if a2 else []
    result["attachment1_vs_attachment2"] = {
        "same_header": a1_header == a2_header,
        "row_counts": [len(a1) - 1, len(a2) - 1],
        "a1_ids": [row[0] for row in a1[1:] if row and row[0] is not None],
        "a2_ids": [row[0] for row in a2[1:] if row and row[0] is not None],
    }

    region_codes: list[Any] = []
    for sheet in by_name["attachment3.xlsx"]["sheets"]:
        for column in sheet["columns"]:
            if column["header"] is not None and "代码" in str(column["header"]):
                region_codes.extend(
                    json.loads(json.dumps(column["examples"]))
                )
    result["attachment3_region_code_checks"] = {
        "sheet_count": len(by_name["attachment3.xlsx"]["sheets"]),
    }

    a4 = by_name["attachment4.xlsx"]
    result["attachment4_template_checks"] = {
        "sheet_count": len(a4["sheets"]),
        "dimensions": [sheet["dimensions"] for sheet in a4["sheets"]],
        "data_row_count": [sheet["data_row_count"] for sheet in a4["sheets"]],
    }
    return result


def render_markdown(audits: list[dict[str, Any]]) -> str:
    lines = ["# 2022 C Data Audit", ""]
    for audit in audits:
        lines.extend(
            [
                f"## {audit['filename']}",
                "",
                f"- SHA256: `{audit['sha256']}`",
                f"- Bytes: `{audit['bytes']}`",
                "",
            ]
        )
        for sheet in audit["sheets"]:
            lines.extend(
                [
                    f"### Sheet: {sheet['name']}",
                    "",
                    f"- Dimension: `{sheet['dimensions']}`",
                    f"- Data rows: `{sheet['data_row_count']}`",
                    f"- Duplicate full rows: `{sheet['duplicate_full_rows']}`",
                    f"- Header: `{sheet['header']}`",
                    "",
                    "| # | Header | Nonempty | Missing | Numeric min | Numeric max | Unique | Top values |",
                    "|---:|---|---:|---:|---:|---:|---:|---|",
                ]
            )
            for column in sheet["columns"]:
                stats = column["range"]
                top = "; ".join(
                    f"{value}:{count}" for value, count in column["top_values"][:6]
                )
                lines.append(
                    "| {index} | {header} | {nonempty} | {missing} | {min} | {max} | {unique} | {top} |".format(
                        index=column["index"],
                        header=str(column["header"]).replace("|", "\\|"),
                        nonempty=column["nonempty"],
                        missing=column["missing"],
                        min=stats["min"],
                        max=stats["max"],
                        unique=stats["unique_count"],
                        top=top.replace("|", "\\|"),
                    )
                )
            lines.append("")
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-root", type=Path)
    parser.add_argument("--output-root", type=Path)
    parser.add_argument("--max-examples", type=int, default=8)
    args = parser.parse_args()

    run_root = Path(__file__).resolve().parents[1]
    default_raw = run_root / "inputs-raw"
    args.raw_root = args.raw_root or default_raw
    args.output_root = args.output_root or (run_root / "work")
    raw_files = [
        args.raw_root / f"attachment{index}.xlsx" for index in range(1, 5)
    ]
    if len(raw_files) != 4:
        raise RuntimeError(f"Expected four attachment workbooks, found {raw_files}")
    missing = [str(path) for path in raw_files if not path.exists()]
    if missing:
        raise RuntimeError(f"Missing ASCII input mirrors: {missing}")
    audits = [audit_workbook(path, args.max_examples) for path in raw_files]
    output = {
        "benchmark_id": "historical_2022_c_buffer_scheduling",
        "run_id": "run-002",
        "audit_scope": [path.name for path in raw_files],
        "workbooks": audits,
    }
    args.output_root.mkdir(parents=True, exist_ok=True)
    (args.output_root / "data-audit.json").write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (args.output_root / "data-audit.md").write_text(
        render_markdown(audits), encoding="utf-8"
    )
    print(json.dumps({"workbooks": [a["filename"] for a in audits]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
