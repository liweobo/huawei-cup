"""Independent audit for run-004 structured-improvement outputs."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
ALLOWED = (
    {0, 1, 2, 3}
    | {lane * 100 + position for lane in range(1, 7) for position in range(1, 11)}
    | {70 + position for position in range(1, 11)}
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit(path: Path) -> dict:
    worksheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    seconds = rows[0][1:]
    matrix = [row[1:] for row in rows[1:]]
    violations = {"code": 0, "position": 0, "receipt": 0}
    for row in matrix:
        for code in row:
            if code is not None and code not in ALLOWED:
                violations["code"] += 1
        if 3 not in row:
            violations["receipt"] += 1
    for second in range(len(seconds)):
        occupied: set[int] = set()
        for row in matrix:
            code = row[second]
            if code is None or code in {0, 1, 2, 3}:
                continue
            if code in occupied:
                violations["position"] += 1
            occupied.add(code)
    return {
        "workbook": path.name,
        "sha256": sha256(path),
        "dimensions": [len(matrix), len(seconds)],
        "violations": violations,
        "total_hard_violations": sum(violations.values()),
        "status": "FEASIBLE" if sum(violations.values()) == 0 else "INFEASIBLE",
    }


def main() -> None:
    report = {"run_id": "run-004-structured-improvement", "results": []}
    for name in ("result41.xlsx", "result42.xlsx"):
        path = ROOT / "outputs" / name
        if path.exists():
            report["results"].append(audit(path))
    report["all_feasible"] = bool(report["results"]) and all(
        item["status"] == "FEASIBLE" for item in report["results"]
    )
    (ROOT / "work" / "feasibility-audit.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
