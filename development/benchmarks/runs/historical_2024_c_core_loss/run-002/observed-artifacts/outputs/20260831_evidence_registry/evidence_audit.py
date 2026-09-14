import argparse
import csv
import hashlib
import json
import math
import os
import platform
import shutil
import subprocess
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import scipy
import sklearn
from scipy.stats import t as tdist
from sklearn.linear_model import HuberRegressor
from sklearn.model_selection import GroupKFold
from sklearn.preprocessing import StandardScaler


WORKSPACE = Path(r"C:\Users\aaa\Desktop\test")
SOURCE_DIR = Path(r"C:\Users\aaa\Desktop\text\pressure-test")
OUTPUT_ROOT = WORKSPACE / "outputs" / "20260831_evidence_registry"
Q1_DIR = WORKSPACE / "outputs" / "20260827_q1_baseline"

TRAIN_PATH = SOURCE_DIR / "\u9644\u4ef6\u4e00\uff08\u8bad\u7ec3\u96c6\uff09.xlsx"
Q1_TEST_PATH = SOURCE_DIR / "\u9644\u4ef6\u4e8c\uff08\u6d4b\u8bd5\u96c6\uff09.xlsx"
Q4_TEST_PATH = SOURCE_DIR / "\u9644\u4ef6\u4e09\uff08\u6d4b\u8bd5\u96c6\uff09.xlsx"
RESULT_TEMPLATE_PATH = SOURCE_DIR / "\u9644\u4ef6\u56db\uff08Excel\u8868\uff09.xlsx"
PROBLEM_PATH = SOURCE_DIR / "\u6570\u636e\u9a71\u52a8\u4e0b\u78c1\u6027\u5143\u4ef6\u7684\u78c1\u82af\u635f\u8017\u5efa\u6a21.docx"
Q1_RESULT_PATH = Q1_DIR / "\u9644\u4ef6\u56db\uff08\u95ee\u9898\u4e00baseline\u7ed3\u679c\uff09.xlsx"

SEED = 20240827
TEMP_VALUES = [25.0, 50.0, 70.0, 90.0]
MODEL_NAMES = ["SE", "linearT", "quadraticT", "interactions"]


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_record(path):
    path = Path(path)
    stat = path.stat()
    return {
        "path": str(path),
        "size_bytes": stat.st_size,
        "mtime_local": datetime.fromtimestamp(stat.st_mtime).isoformat(),
        "sha256": sha256_file(path),
    }


def run_command(command, cwd, stdout_path, stderr_path):
    started = datetime.now()
    started_clock = time.perf_counter()
    process = subprocess.run(
        command,
        cwd=cwd,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    elapsed = time.perf_counter() - started_clock
    stdout_path.write_text(process.stdout, encoding="utf-8")
    stderr_path.write_text(process.stderr, encoding="utf-8")
    record = {
        "command": command,
        "cwd": str(cwd),
        "started_local": started.isoformat(),
        "finished_local": datetime.now().isoformat(),
        "elapsed_seconds": elapsed,
        "exit_code": process.returncode,
        "stdout": str(stdout_path),
        "stderr": str(stderr_path),
    }
    if process.returncode != 0:
        raise RuntimeError(json.dumps(record, ensure_ascii=False, indent=2))
    return record


def metrics(P, y, idx, prediction):
    true = P[idx]
    error = prediction - true
    percentage_error = np.abs(error / true)
    return {
        "n": int(len(idx)),
        "MAPE": float(percentage_error.mean()),
        "MdAPE": float(np.median(percentage_error)),
        "RMSE_log": float(
            np.sqrt(np.mean((np.log(prediction) - y[idx]) ** 2))
        ),
        "NRMSE_mean": float(np.sqrt(np.mean(error * error)) / np.mean(true)),
    }


def build_design(T, f, B, model):
    tc = (T - 25.0) / 65.0
    lf = np.log(f)
    lb = np.log(B)
    if model == "SE":
        return np.column_stack([np.ones(len(T)), lf, lb])
    if model == "linearT":
        return np.column_stack([np.ones(len(T)), lf, lb, tc])
    if model == "quadraticT":
        return np.column_stack([np.ones(len(T)), lf, lb, tc, tc * tc])
    lfc = lf - lf.mean()
    lbc = lb - lb.mean()
    return np.column_stack(
        [np.ones(len(T)), lfc, lbc, tc, tc * lfc, tc * lbc]
    )


def fit_predict(P, y, X, train, test):
    beta = np.linalg.lstsq(X[train], y[train], rcond=None)[0]
    residual = y[train] - X[train] @ beta
    smear = float(np.mean(np.exp(residual)))
    prediction = np.exp(X[test] @ beta) * smear
    return prediction, beta, smear


def compare_models(T, f, P, B, train, test):
    y = np.log(P)
    result = {}
    for model in MODEL_NAMES:
        X = build_design(T, f, B, model)
        prediction, beta, smear = fit_predict(P, y, X, train, test)
        result[model] = {
            **metrics(P, y, test, prediction),
            "beta": [float(value) for value in beta],
            "smear": smear,
        }
    return result


def summarize_repeats(repeats):
    result = {}
    for model in MODEL_NAMES:
        result[model] = {}
        for metric in ["MAPE", "MdAPE", "RMSE_log", "NRMSE_mean"]:
            values = np.asarray([run[model][metric] for run in repeats])
            result[model][metric] = {
                "mean": float(values.mean()),
                "sd": float(values.std(ddof=1)),
                "min": float(values.min()),
                "max": float(values.max()),
            }
    return result


def read_training_data():
    workbook = openpyxl.load_workbook(TRAIN_PATH, read_only=True, data_only=True)
    q2_rows = []
    q3_counts_raw = Counter()
    q3_counts_clean = Counter()
    duplicate_rows = []
    seen_by_sheet = defaultdict(dict)
    total_raw = 0
    total_clean = 0

    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            total_raw += 1
            temperature = float(row[0])
            frequency = float(row[1])
            loss = float(row[2])
            waveform = str(row[3])
            wave = np.asarray(row[4:], dtype=np.float64)
            key = (sheet.title, int(temperature), waveform)
            q3_counts_raw[key] += 1

            digest = hashlib.sha256()
            digest.update(
                (f"{temperature}|{frequency}|{loss}|{waveform}|").encode("utf-8")
            )
            digest.update(wave.tobytes())
            row_hash = digest.hexdigest()
            if row_hash in seen_by_sheet[sheet.title]:
                duplicate_rows.append(
                    {
                        "sheet": sheet.title,
                        "row": row_number,
                        "duplicates_row": seen_by_sheet[sheet.title][row_hash],
                        "row_sha256": row_hash,
                    }
                )
                continue
            seen_by_sheet[sheet.title][row_hash] = row_number
            q3_counts_clean[key] += 1
            total_clean += 1

            if sheet.title == "\u6750\u65991" and waveform == "\u6b63\u5f26\u6ce2":
                q2_rows.append(
                    (
                        row_number,
                        temperature,
                        frequency,
                        loss,
                        float((wave.max() - wave.min()) / 2.0),
                        float(np.max(np.abs(wave))),
                    )
                )
    workbook.close()
    return {
        "q2_rows": np.asarray(q2_rows, dtype=float),
        "q3_counts_raw": q3_counts_raw,
        "q3_counts_clean": q3_counts_clean,
        "total_raw": total_raw,
        "total_clean": total_clean,
        "duplicate_rows": duplicate_rows,
    }


def q2_audit(q2_rows):
    row_number, T, f, P, B_half, B_abs = q2_rows.T
    y = np.log(P)
    all_indices = np.arange(len(P))

    random_runs = []
    for repeat in range(50):
        rng = np.random.default_rng(SEED + repeat)
        train = []
        test = []
        for temperature in TEMP_VALUES:
            indices = np.where(T == temperature)[0].copy()
            rng.shuffle(indices)
            n_test = int(round(0.2 * len(indices)))
            test.extend(indices[:n_test])
            train.extend(indices[n_test:])
        random_runs.append(
            compare_models(
                T, f, P, B_half, np.asarray(train), np.asarray(test)
            )
        )

    leave_one_temperature_out = {}
    for temperature in TEMP_VALUES:
        test = np.where(T == temperature)[0]
        train = np.where(T != temperature)[0]
        leave_one_temperature_out[str(int(temperature))] = compare_models(
            T, f, P, B_half, train, test
        )

    split = int(round(0.8 * len(P)))
    order_train = all_indices[:split]
    order_test = all_indices[split:]
    order_stress = {
        "meta": {
            "train_temperature_counts": {
                str(int(value)): int(np.sum(T[:split] == value))
                for value in TEMP_VALUES
            },
            "test_temperature_counts": {
                str(int(value)): int(np.sum(T[split:] == value))
                for value in TEMP_VALUES
            },
            "test_frequency_range": [
                float(f[split:].min()),
                float(f[split:].max()),
            ],
            "test_Bm_range": [
                float(B_half[split:].min()),
                float(B_half[split:].max()),
            ],
        },
        "metrics": compare_models(T, f, P, B_half, order_train, order_test),
    }

    group_runs = []
    group_cv = GroupKFold(n_splits=5)
    for train, test in group_cv.split(all_indices, groups=f.astype(int)):
        group_runs.append(compare_models(T, f, P, B_half, train, test))

    full_fit = {}
    for label, B in [("half_range", B_half), ("max_abs", B_abs)]:
        full_fit[label] = compare_models(T, f, P, B, all_indices, all_indices)

    # Robust uncertainty and model-form sensitivity for the linear temperature term.
    tc = (T - 25.0) / 65.0
    lf = np.log(f)
    lb = np.log(B_half)
    X = np.column_stack([np.ones(len(P)), lf, lb, tc])
    beta = np.linalg.lstsq(X, y, rcond=None)[0]
    residual = y - X @ beta
    inverse = np.linalg.inv(X.T @ X)
    leverage = np.sum((X @ inverse) * X, axis=1)
    meat = X.T @ (((residual / (1.0 - leverage)) ** 2)[:, None] * X)
    covariance_hc3 = inverse @ meat @ inverse
    se_hc3 = np.sqrt(np.diag(covariance_hc3))
    critical = tdist.ppf(0.975, len(P) - X.shape[1])
    ci_hc3 = np.column_stack(
        [beta - critical * se_hc3, beta + critical * se_hc3]
    )

    groups = f.astype(int)
    unique_groups = np.unique(groups)
    cluster_meat = np.zeros((X.shape[1], X.shape[1]))
    for group in unique_groups:
        selected = groups == group
        score = X[selected].T @ residual[selected]
        cluster_meat += np.outer(score, score)
    G = len(unique_groups)
    N = len(P)
    K = X.shape[1]
    covariance_cluster = (
        inverse
        @ cluster_meat
        @ inverse
        * (G / (G - 1))
        * ((N - 1) / (N - K))
    )
    se_cluster = np.sqrt(np.diag(covariance_cluster))
    critical_cluster = tdist.ppf(0.975, G - 1)
    ci_cluster = np.column_stack(
        [
            beta - critical_cluster * se_cluster,
            beta + critical_cluster * se_cluster,
        ]
    )

    keep = np.abs(residual) <= np.quantile(np.abs(residual), 0.99)
    beta_trimmed = np.linalg.lstsq(X[keep], y[keep], rcond=None)[0]
    scaler = StandardScaler().fit(np.column_stack([lf, lb, tc]))
    standardized = scaler.transform(np.column_stack([lf, lb, tc]))
    huber = HuberRegressor(
        alpha=0.0, epsilon=1.35, max_iter=1000
    ).fit(standardized, y)
    beta_huber = huber.coef_ / scaler.scale_

    X_abs = np.column_stack([np.ones(len(P)), lf, np.log(B_abs), tc])
    beta_abs = np.linalg.lstsq(X_abs, y, rcond=None)[0]

    dummies = np.column_stack(
        [(T == value).astype(float) for value in [50.0, 70.0, 90.0]]
    )
    X_category = np.column_stack([np.ones(len(P)), lf, lb, dummies])
    beta_category = np.linalg.lstsq(X_category, y, rcond=None)[0]

    relative_B_difference = np.abs(B_abs - B_half) / B_half
    return {
        "status": "verified_real_run",
        "seed": SEED,
        "sample_count": int(len(P)),
        "temperature_counts": {
            str(int(value)): int(np.sum(T == value)) for value in TEMP_VALUES
        },
        "frequency_range": [float(f.min()), float(f.max())],
        "Bm_half_range": [float(B_half.min()), float(B_half.max())],
        "loss_range": [float(P.min()), float(P.max())],
        "Bm_definition_relative_difference": {
            "median": float(np.median(relative_B_difference)),
            "max": float(relative_B_difference.max()),
        },
        "random_stratified_50_repeats": summarize_repeats(random_runs),
        "leave_one_temperature_out": leave_one_temperature_out,
        "excel_order_last20_stress_not_time_split": order_stress,
        "unseen_frequency_group_cv": summarize_repeats(group_runs),
        "full_fit": full_fit,
        "linear_temperature_inference": {
            "beta": [float(value) for value in beta],
            "gamma_per_C": float(beta[3] / 65.0),
            "factor_90_vs_25": float(np.exp(beta[3])),
            "HC3_95CI_gamma_per_C": [float(value) for value in ci_hc3[3] / 65.0],
            "frequency_cluster_95CI_gamma_per_C": [
                float(value) for value in ci_cluster[3] / 65.0
            ],
            "frequency_cluster_count": int(G),
        },
        "sensitivity_gamma_per_C": {
            "max_abs_Bm": float(beta_abs[3] / 65.0),
            "trim_largest_1pct_log_residuals": float(beta_trimmed[3] / 65.0),
            "Huber": float(beta_huber[2] / 65.0),
        },
        "categorical_temperature_factors_vs_25": {
            "50": float(np.exp(beta_category[3])),
            "70": float(np.exp(beta_category[4])),
            "90": float(np.exp(beta_category[5])),
        },
    }


def q3_audit(training_record):
    raw_counts = training_record["q3_counts_raw"]
    clean_counts = training_record["q3_counts_clean"]
    materials = sorted({key[0] for key in raw_counts})
    temperatures = sorted({key[1] for key in raw_counts})
    waveforms = sorted({key[2] for key in raw_counts})
    expected_cells = len(materials) * len(temperatures) * len(waveforms)
    return {
        "status": "data_design_verified_model_not_run",
        "raw_total_rows": training_record["total_raw"],
        "clean_total_rows": training_record["total_clean"],
        "duplicate_rows": training_record["duplicate_rows"],
        "materials": materials,
        "temperatures": temperatures,
        "waveforms": waveforms,
        "expected_factor_cells": expected_cells,
        "observed_factor_cells_raw": len(raw_counts),
        "all_factor_cells_present": len(raw_counts) == expected_cells,
        "raw_cell_count_minmax": [min(raw_counts.values()), max(raw_counts.values())],
        "clean_cell_count_minmax": [
            min(clean_counts.values()),
            max(clean_counts.values()),
        ],
        "raw_cell_counts": {
            f"{material}|{temperature}|{waveform}": raw_counts[
                (material, temperature, waveform)
            ]
            for material in materials
            for temperature in temperatures
            for waveform in waveforms
        },
        "interaction_effects": {
            "status": "not_run",
            "note": "Only the analysis protocol has been defined; no interaction-effect number is paper-ready yet.",
        },
    }


def read_sheet_rows(path, sheet_name, max_columns=5, data_only=True):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    sheet = workbook[sheet_name]
    rows = list(
        sheet.iter_rows(
            min_row=2,
            max_row=401,
            min_col=1,
            max_col=max_columns,
            values_only=True,
        )
    )
    workbook.close()
    return rows


def q4_alignment_audit():
    test_values = read_sheet_rows(Q4_TEST_PATH, "\u6d4b\u8bd5\u96c6", 5, True)
    template_values = read_sheet_rows(RESULT_TEMPLATE_PATH, "Sheet1", 3, True)
    template_formulas = read_sheet_rows(RESULT_TEMPLATE_PATH, "Sheet1", 3, False)
    q1_values = read_sheet_rows(Q1_RESULT_PATH, "Sheet1", 3, True)

    ids_test = [row[0] for row in test_values]
    ids_template = [row[0] for row in template_values]
    ids_q1 = [row[0] for row in q1_values]
    expected = list(range(1, 401))
    special = [16, 76, 98, 126, 168, 230, 271, 338, 348, 379]
    return {
        "status": "alignment_verified_predictions_not_run_or_written",
        "attachment3_rows": len(test_values),
        "attachment4_rows": len(template_values),
        "attachment3_ids_are_1_to_400": ids_test == expected,
        "attachment4_ids_are_1_to_400": ids_template == expected,
        "q1_result_ids_are_1_to_400": ids_q1 == expected,
        "rowwise_id_match": ids_test == ids_template == ids_q1,
        "attachment4_original_col2_nonblank": int(
            sum(row[1] is not None for row in template_values)
        ),
        "attachment4_original_col3_nonblank": int(
            sum(row[2] is not None for row in template_values)
        ),
        "q1_result_col2_nonblank": int(
            sum(row[1] is not None for row in q1_values)
        ),
        "q1_result_col3_nonblank": int(
            sum(row[2] is not None for row in q1_values)
        ),
        "sequence_formula_first5": [row[0] for row in template_formulas[:5]],
        "sequence_formula_last5": [row[0] for row in template_formulas[-5:]],
        "prediction_target_range": "Sheet1!C2:C401",
        "sample_to_excel_row_rule": "excel_row = sample_id + 1",
        "special_mapping": [
            {
                "sample_id": sample_id,
                "excel_row": sample_id + 1,
                "target_cell": f"C{sample_id + 1}",
                "temperature": test_values[sample_id - 1][1],
                "frequency": test_values[sample_id - 1][2],
                "material": test_values[sample_id - 1][3],
                "waveform": test_values[sample_id - 1][4],
            }
            for sample_id in special
        ],
        "q4_predictions": {
            "status": "not_run",
            "note": "No question-four loss predictions are claimed or written yet.",
        },
    }


def get_path(record, path):
    current = record
    for item in path.split("/"):
        if item == "":
            continue
        current = current[int(item)] if isinstance(current, list) else current[item]
    return current


def claim(claim_id, section, description, record_name, path, expected, tolerance=1e-12):
    return {
        "claim_id": claim_id,
        "section": section,
        "description": description,
        "record": record_name,
        "json_pointer": "/" + path,
        "expected_from_prior_analysis": expected,
        "tolerance": tolerance,
    }


def build_claims(records):
    claims = [
        claim("Q1-001", "Q1", "Clean training rows", "q1_baseline", "training_rows_after_duplicate_removal", 12399),
        claim("Q1-002", "Q1", "Gaussian NB holdout accuracy", "q1_baseline", "candidate_validation_metrics/gaussian_naive_bayes/accuracy", 1.0),
        claim("Q1-003", "Q1", "Gaussian NB holdout macro F1", "q1_baseline", "candidate_validation_metrics/gaussian_naive_bayes/macro_f1", 1.0),
        claim("Q1-004", "Q1", "Gaussian NB leave-one-material accuracy", "q1_baseline", "leave_one_material_out_metrics/gaussian_naive_bayes/aggregate/accuracy", 1.0),
        claim("Q1-005", "Q1", "RBF-SVM holdout accuracy", "q1_comparison", "validation_metrics/rbf_svm/accuracy", 1.0),
        claim("Q1-006", "Q1", "Random forest holdout accuracy", "q1_comparison", "validation_metrics/random_forest/accuracy", 1.0),
        claim("Q1-007", "Q1", "RBF-SVM leave-one-material accuracy", "q1_comparison", "leave_one_material_out_metrics/rbf_svm/aggregate/accuracy", 1.0),
        claim("Q1-008", "Q1", "Random forest leave-one-material accuracy", "q1_comparison", "leave_one_material_out_metrics/random_forest/aggregate/accuracy", 1.0),
        claim("Q1-009", "Q1", "Test sine count", "q1_comparison", "test_prediction_counts/rbf_svm/\u6b63\u5f26\u6ce2", 20),
        claim("Q1-010", "Q1", "Test triangle count", "q1_comparison", "test_prediction_counts/rbf_svm/\u4e09\u89d2\u6ce2", 44),
        claim("Q1-011", "Q1", "Test trapezoid count", "q1_comparison", "test_prediction_counts/rbf_svm/\u68af\u5f62\u6ce2", 16),
        claim("Q1-012", "Q1", "Candidate agreement on 80 test samples", "q1_comparison", "test_agreement/rbf_svm/with_other_candidate", 1.0),
        claim("Q2-001", "Q2", "Material1 sine sample count", "q2", "sample_count", 1067),
        claim("Q2-002", "Q2", "SE repeated-split MAPE", "q2", "random_stratified_50_repeats/SE/MAPE/mean", 0.3291511436035553),
        claim("Q2-003", "Q2", "Linear temperature repeated-split MAPE", "q2", "random_stratified_50_repeats/linearT/MAPE/mean", 0.1687979253107534),
        claim("Q2-004", "Q2", "Quadratic temperature repeated-split MAPE", "q2", "random_stratified_50_repeats/quadraticT/MAPE/mean", 0.16164447236535204),
        claim("Q2-005", "Q2", "Interaction model repeated-split MAPE", "q2", "random_stratified_50_repeats/interactions/MAPE/mean", 0.12926035908793337),
        claim("Q2-006", "Q2", "SE unseen-frequency MAPE", "q2", "unseen_frequency_group_cv/SE/MAPE/mean", 0.32473914538513265),
        claim("Q2-007", "Q2", "Linear temperature unseen-frequency MAPE", "q2", "unseen_frequency_group_cv/linearT/MAPE/mean", 0.17513120856552475),
        claim("Q2-008", "Q2", "Quadratic temperature unseen-frequency MAPE", "q2", "unseen_frequency_group_cv/quadraticT/MAPE/mean", 0.16833749139761212),
        claim("Q2-009", "Q2", "Interaction model unseen-frequency MAPE", "q2", "unseen_frequency_group_cv/interactions/MAPE/mean", 0.137945023104263),
        claim("Q2-010", "Q2", "Linear gamma per C", "q2", "linear_temperature_inference/gamma_per_C", -0.011857064562093735),
        claim("Q2-011", "Q2", "Linear factor 90C versus 25C", "q2", "linear_temperature_inference/factor_90_vs_25", 0.4626848174579582),
        claim("Q2-012", "Q2", "Bm definition median relative difference", "q2", "Bm_definition_relative_difference/median", 0.0015641285638175364),
        claim("Q2-013", "Q2", "Bm definition max relative difference", "q2", "Bm_definition_relative_difference/max", 0.010712499522593767),
        claim("Q3-001", "Q3", "Raw training rows", "q3", "raw_total_rows", 12400),
        claim("Q3-002", "Q3", "Complete factor cells", "q3", "expected_factor_cells", 48),
        claim("Q3-003", "Q3", "Minimum raw cell count", "q3", "raw_cell_count_minmax/0", 63),
        claim("Q3-004", "Q3", "Maximum raw cell count", "q3", "raw_cell_count_minmax/1", 445),
        claim("Q4-001", "Q4", "Attachment3 sample rows", "q4_alignment", "attachment3_rows", 400),
        claim("Q4-002", "Q4", "Attachment4 result rows", "q4_alignment", "attachment4_rows", 400),
        claim("Q4-003", "Q4", "Existing Q1 results in column B", "q4_alignment", "q1_result_col2_nonblank", 80),
        claim("Q4-004", "Q4", "Q4 prediction cells currently populated", "q4_alignment", "q1_result_col3_nonblank", 0),
    ]
    for item in claims:
        actual = get_path(records[item["record"]], item["json_pointer"])
        expected = item["expected_from_prior_analysis"]
        if isinstance(actual, (float, int)) and isinstance(expected, (float, int)):
            matches = math.isclose(
                float(actual), float(expected), rel_tol=item["tolerance"], abs_tol=item["tolerance"]
            )
        else:
            matches = actual == expected
        item["actual_from_current_run"] = actual
        item["matches"] = matches
        item["status"] = "verified" if matches else "mismatch"
    return claims


def flatten_numbers(value, pointer=""):
    rows = []
    if isinstance(value, dict):
        for key, child in value.items():
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            rows.extend(flatten_numbers(child, f"{pointer}/{escaped}"))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            rows.extend(flatten_numbers(child, f"{pointer}/{index}"))
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        rows.append((pointer or "/", value))
    return rows


def write_csv(path, fieldnames, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", default=datetime.now().strftime("%Y%m%dT%H%M%S"))
    parser.add_argument("--skip-q1-rerun", action="store_true")
    args = parser.parse_args()

    run_dir = OUTPUT_ROOT / args.run_id
    run_dir.mkdir(parents=True, exist_ok=False)
    q1_snapshot_dir = run_dir / "q1_snapshot"
    q1_snapshot_dir.mkdir()

    run_started = datetime.now()
    inputs = {
        "training": file_record(TRAIN_PATH),
        "q1_test": file_record(Q1_TEST_PATH),
        "q4_test": file_record(Q4_TEST_PATH),
        "result_template": file_record(RESULT_TEMPLATE_PATH),
        "problem_statement": file_record(PROBLEM_PATH),
    }

    q1_run_records = []
    if not args.skip_q1_rerun:
        q1_run_records.append(
            run_command(
                [sys.executable, str(Q1_DIR / "q1_baseline.py")],
                Q1_DIR,
                run_dir / "q1_baseline.stdout.log",
                run_dir / "q1_baseline.stderr.log",
            )
        )
        q1_run_records.append(
            run_command(
                [sys.executable, str(Q1_DIR / "compare_models.py")],
                Q1_DIR,
                run_dir / "q1_compare.stdout.log",
                run_dir / "q1_compare.stderr.log",
            )
        )

    q1_files = [
        "summary.json",
        "predictions.csv",
        "model_comparison.json",
        "model_comparison_predictions.csv",
        Q1_RESULT_PATH.name,
    ]
    for filename in q1_files:
        shutil.copy2(Q1_DIR / filename, q1_snapshot_dir / filename)

    with open(Q1_DIR / "summary.json", encoding="utf-8") as handle:
        q1_baseline = json.load(handle)
    with open(Q1_DIR / "model_comparison.json", encoding="utf-8") as handle:
        q1_comparison = json.load(handle)

    training_record = read_training_data()
    q2 = q2_audit(training_record["q2_rows"])
    q3 = q3_audit(training_record)
    q4_alignment = q4_alignment_audit()
    records = {
        "q1_baseline": q1_baseline,
        "q1_comparison": q1_comparison,
        "q2": q2,
        "q3": q3,
        "q4_alignment": q4_alignment,
    }
    claims = build_claims(records)

    q2_path = run_dir / "q2_validation.json"
    q3_path = run_dir / "q3_data_design_audit.json"
    q4_path = run_dir / "q4_alignment_audit.json"
    q2_path.write_text(json.dumps(q2, ensure_ascii=False, indent=2), encoding="utf-8")
    q3_path.write_text(json.dumps(q3, ensure_ascii=False, indent=2), encoding="utf-8")
    q4_path.write_text(
        json.dumps(q4_alignment, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    claim_rows = []
    for item in claims:
        claim_rows.append(
            {
                "claim_id": item["claim_id"],
                "section": item["section"],
                "description": item["description"],
                "status": item["status"],
                "expected_from_prior_analysis": item["expected_from_prior_analysis"],
                "actual_from_current_run": item["actual_from_current_run"],
                "record": item["record"],
                "json_pointer": item["json_pointer"],
            }
        )
    claims_path = run_dir / "paper_claims_registry.csv"
    write_csv(
        claims_path,
        [
            "claim_id",
            "section",
            "description",
            "status",
            "expected_from_prior_analysis",
            "actual_from_current_run",
            "record",
            "json_pointer",
        ],
        claim_rows,
    )

    numeric_rows = []
    record_files = {
        "q1_baseline": q1_snapshot_dir / "summary.json",
        "q1_comparison": q1_snapshot_dir / "model_comparison.json",
        "q2": q2_path,
        "q3": q3_path,
        "q4_alignment": q4_path,
    }
    for record_name, record in records.items():
        for pointer, value in flatten_numbers(record):
            numeric_rows.append(
                {
                    "record": record_name,
                    "json_pointer": pointer,
                    "value": value,
                    "evidence_file": str(record_files[record_name]),
                    "status": "real_run_or_direct_audit",
                }
            )
    numeric_path = run_dir / "all_numeric_evidence.csv"
    write_csv(
        numeric_path,
        ["record", "json_pointer", "value", "evidence_file", "status"],
        numeric_rows,
    )

    supporting_files = [
        q2_path,
        q3_path,
        q4_path,
        claims_path,
        numeric_path,
        *list(q1_snapshot_dir.iterdir()),
    ]
    output_records = {path.name: file_record(path) for path in supporting_files}
    script_records = {
        "evidence_audit.py": file_record(Path(__file__)),
        "q1_baseline.py": file_record(Q1_DIR / "q1_baseline.py"),
        "compare_models.py": file_record(Q1_DIR / "compare_models.py"),
    }

    all_claims_verified = all(item["matches"] for item in claims)
    manifest = {
        "run_id": args.run_id,
        "status": "verified" if all_claims_verified else "mismatch_detected",
        "started_local": run_started.isoformat(),
        "finished_local": datetime.now().isoformat(),
        "timezone": "Asia/Shanghai",
        "environment": {
            "python": sys.version,
            "python_executable": sys.executable,
            "platform": platform.platform(),
            "numpy": np.__version__,
            "openpyxl": openpyxl.__version__,
            "scipy": scipy.__version__,
            "sklearn": sklearn.__version__,
        },
        "inputs": inputs,
        "scripts": script_records,
        "q1_run_commands": q1_run_records,
        "records": {name: str(path) for name, path in record_files.items()},
        "outputs": output_records,
        "claims_verified": sum(item["matches"] for item in claims),
        "claims_total": len(claims),
        "all_claims_verified": all_claims_verified,
        "paper_readiness": {
            "Q1": "verified_real_run",
            "Q2": "verified_real_run",
            "Q3_data_design": "verified_direct_audit",
            "Q3_interaction_effects": "not_run_not_paper_ready",
            "Q4_alignment": "verified_direct_audit",
            "Q4_predictions": "not_run_not_written_not_paper_ready",
        },
    }
    manifest_path = run_dir / "evidence_manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    mismatches = [item for item in claims if not item["matches"]]
    summary_lines = [
        f"# Evidence Run {args.run_id}",
        "",
        f"- Status: {'VERIFIED' if all_claims_verified else 'MISMATCH DETECTED'}",
        f"- Curated numeric claims verified: {len(claims) - len(mismatches)}/{len(claims)}",
        f"- Numeric leaves registered: {len(numeric_rows)}",
        f"- Training input SHA256: `{inputs['training']['sha256']}`",
        "",
        "## Paper Readiness",
        "",
        "- Q1 metrics and predictions: verified real run.",
        "- Q2 validation, generalization, and sensitivity: verified real run.",
        "- Q3 factor-cell counts: verified direct audit; interaction effects not run.",
        "- Q4 row/column alignment: verified direct audit; predictions not run or written.",
        "",
        "## Mismatches",
        "",
    ]
    if mismatches:
        summary_lines.extend(
            [
                f"- {item['claim_id']}: expected {item['expected_from_prior_analysis']}, got {item['actual_from_current_run']}"
                for item in mismatches
            ]
        )
    else:
        summary_lines.append("- None.")
    summary_path = run_dir / "README.md"
    summary_path.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    checksum_paths = [manifest_path, summary_path, *supporting_files]
    checksum_lines = [
        f"{sha256_file(path)}  {path.relative_to(run_dir)}" for path in checksum_paths
    ]
    (run_dir / "checksums.sha256").write_text(
        "\n".join(checksum_lines) + "\n", encoding="utf-8"
    )

    print(
        json.dumps(
            {
                "run_dir": str(run_dir),
                "status": manifest["status"],
                "claims_verified": manifest["claims_verified"],
                "claims_total": manifest["claims_total"],
                "numeric_evidence_count": len(numeric_rows),
                "paper_readiness": manifest["paper_readiness"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
