import argparse
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

import openpyxl


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("run_dir")
    args = parser.parse_args()
    run_dir = Path(args.run_dir)

    manifest_path = run_dir / "evidence_manifest.json"
    claims_path = run_dir / "paper_claims_registry.csv"
    predictions_path = run_dir / "q1_snapshot" / "predictions.csv"
    workbook_path = (
        run_dir
        / "q1_snapshot"
        / "\u9644\u4ef6\u56db\uff08\u95ee\u9898\u4e00baseline\u7ed3\u679c\uff09.xlsx"
    )

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    with open(claims_path, encoding="utf-8-sig") as handle:
        claims = list(csv.DictReader(handle))
    with open(predictions_path, encoding="utf-8-sig") as handle:
        predictions = list(csv.DictReader(handle))

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    mismatches = []
    for excel_row, prediction in enumerate(predictions, start=2):
        expected_id = int(prediction["\u5e8f\u53f7"])
        expected_class = int(prediction["\u5206\u7c7b\u7f16\u53f7"])
        actual_id = sheet.cell(excel_row, 1).value
        actual_class = sheet.cell(excel_row, 2).value
        if actual_id != expected_id or actual_class != expected_class:
            mismatches.append(
                {
                    "excel_row": excel_row,
                    "expected_id": expected_id,
                    "actual_id": actual_id,
                    "expected_class": expected_class,
                    "actual_class": actual_class,
                }
            )
    workbook.close()

    checksum_results = []
    with open(run_dir / "checksums.sha256", encoding="utf-8") as handle:
        for line in handle:
            expected_hash, relative_path = line.rstrip("\n").split("  ", 1)
            target = run_dir / relative_path
            actual_hash = sha256_file(target)
            checksum_results.append(
                {
                    "path": relative_path,
                    "expected_sha256": expected_hash,
                    "actual_sha256": actual_hash,
                    "matches": expected_hash == actual_hash,
                }
            )

    report = {
        "verified_local": datetime.now().isoformat(),
        "run_id": manifest["run_id"],
        "manifest_status": manifest["status"],
        "claims_verified": sum(row["status"] == "verified" for row in claims),
        "claims_total": len(claims),
        "all_claims_verified": all(row["status"] == "verified" for row in claims),
        "q1_prediction_rows": len(predictions),
        "q1_prediction_xlsx_mismatches": mismatches,
        "q1_predictions_match_xlsx": not mismatches and len(predictions) == 80,
        "checksum_entries": len(checksum_results),
        "checksums_all_match": all(row["matches"] for row in checksum_results),
        "checksum_results": checksum_results,
        "paper_readiness": manifest["paper_readiness"],
        "source_hashes": {
            "manifest": sha256_file(manifest_path),
            "claims_registry": sha256_file(claims_path),
            "q1_predictions": sha256_file(predictions_path),
            "q1_result_workbook": sha256_file(workbook_path),
            "verifier_script": sha256_file(Path(__file__)),
        },
    }
    report["verified"] = (
        report["manifest_status"] == "verified"
        and report["all_claims_verified"]
        and report["q1_predictions_match_xlsx"]
        and report["checksums_all_match"]
    )

    report_path = run_dir / "post_run_verification.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    report_hash = sha256_file(report_path)
    (run_dir / "post_run_verification.sha256").write_text(
        f"{report_hash}  post_run_verification.json\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "verified": report["verified"],
                "claims": [report["claims_verified"], report["claims_total"]],
                "q1_prediction_xlsx_mismatches": len(mismatches),
                "checksums": [
                    sum(row["matches"] for row in checksum_results),
                    len(checksum_results),
                ],
                "report": str(report_path),
                "report_sha256": report_hash,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
