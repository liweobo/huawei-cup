import csv
import json
import math
import os
import random
from collections import Counter

import numpy as np
import openpyxl


TRAIN_PATH = r"C:\Users\aaa\Desktop\text\pressure-test\附件一（训练集）.xlsx"
TEST_PATH = r"C:\Users\aaa\Desktop\text\pressure-test\附件二（测试集）.xlsx"
OUT_DIR = os.path.dirname(os.path.abspath(__file__))
SEED = 20240827
LABELS = ["正弦波", "三角波", "梯形波"]
LABEL_TO_INT = {label: i + 1 for i, label in enumerate(LABELS)}


def shape_features(values):
    """Extract a small, phase-robust feature vector from one waveform."""
    x = np.asarray(values, dtype=np.float32)
    x = x - np.mean(x)
    amplitude = float(np.max(x) - np.min(x))
    x = x / (amplitude if amplitude > 1e-12 else 1.0)

    spectrum = np.abs(np.fft.rfft(x))[1:17].astype(np.float32)
    spectrum = spectrum / max(float(spectrum[0]), 1e-8)

    derivative = np.roll(x, -1) - x
    abs_derivative = np.abs(derivative)
    level_quantiles = np.quantile(
        x, [0.05, 0.10, 0.25, 0.50, 0.75, 0.90, 0.95]
    ).astype(np.float32)
    slope_quantiles = np.quantile(
        abs_derivative, [0.25, 0.50, 0.75, 0.90, 0.95]
    ).astype(np.float32)
    other = np.asarray(
        [
            np.sqrt(np.mean(x * x)),
            np.mean(abs_derivative),
            np.std(derivative),
            np.max(abs_derivative),
            np.mean(abs_derivative < 0.002),
            np.mean(abs_derivative < 0.005),
            np.mean(np.abs(x) < 0.05),
            np.mean(x > 0.40),
            np.mean(x < -0.40),
        ],
        dtype=np.float32,
    )
    return np.concatenate([spectrum, level_quantiles, slope_quantiles, other])


def read_train():
    workbook = openpyxl.load_workbook(TRAIN_PATH, read_only=True, data_only=True)
    records = []
    seen = set()
    duplicate_removed = []
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            key = (sheet.title, tuple(row[:4]), tuple(float(v) for v in row[4:]))
            if key in seen:
                duplicate_removed.append({"sheet": sheet.title, "row": row_number})
                continue
            seen.add(key)
            records.append(
                (sheet.title, LABEL_TO_INT[row[3]], shape_features(row[4:]))
            )
    workbook.close()
    X = np.vstack([record[2] for record in records])
    y = np.asarray([record[1] for record in records], dtype=np.int64)
    groups = np.asarray([record[0] for record in records], dtype=object)
    return X, y, groups, duplicate_removed


def read_test():
    workbook = openpyxl.load_workbook(TEST_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    ids, features = [], []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        ids.append(int(row[0]))
        features.append(shape_features(row[4:]))
    workbook.close()
    return np.asarray(ids), np.vstack(features)


def fit_nearest_centroid(X, y):
    mean = X.mean(axis=0)
    scale = X.std(axis=0)
    scale[scale < 1e-8] = 1.0
    Z = (X - mean) / scale
    centroids = np.vstack([Z[y == c].mean(axis=0) for c in (1, 2, 3)])
    return mean, scale, centroids


def predict_nearest_centroid(X, parameters):
    mean, scale, centroids = parameters
    Z = (X - mean) / scale
    distances = ((Z[:, None, :] - centroids[None, :, :]) ** 2).sum(axis=2)
    return np.argmin(distances, axis=1) + 1, distances


def fit_gaussian_nb(X, y):
    means, variances, log_priors = [], [], []
    for category in (1, 2, 3):
        values = X[y == category].astype(np.float64)
        means.append(values.mean(axis=0))
        variance = values.var(axis=0)
        variance += max(float(np.mean(variance)) * 1e-6, 1e-10)
        variances.append(variance)
        log_priors.append(math.log(len(values) / len(X)))
    return np.vstack(means), np.vstack(variances), np.asarray(log_priors)


def predict_gaussian_nb(X, parameters):
    means, variances, log_priors = parameters
    values = X.astype(np.float64)
    scores = []
    for k in range(3):
        score = -0.5 * np.sum(
            np.log(2 * np.pi * variances[k])
            + (values - means[k]) ** 2 / variances[k],
            axis=1,
        )
        scores.append(score + log_priors[k])
    scores = np.vstack(scores).T
    return np.argmax(scores, axis=1) + 1, -scores


def fit_predict(model, X_train, y_train, X_valid):
    if model == "gaussian_naive_bayes":
        return predict_gaussian_nb(X_valid, fit_gaussian_nb(X_train, y_train))
    return predict_nearest_centroid(
        X_valid, fit_nearest_centroid(X_train, y_train)
    )


def stratified_split(y, fraction=0.2):
    rng = random.Random(SEED)
    train_indices, valid_indices = [], []
    for category in (1, 2, 3):
        indices = np.flatnonzero(y == category).tolist()
        rng.shuffle(indices)
        n_valid = int(round(len(indices) * fraction))
        valid_indices.extend(indices[:n_valid])
        train_indices.extend(indices[n_valid:])
    return np.asarray(train_indices), np.asarray(valid_indices)


def metrics(y_true, y_pred):
    matrix = [
        [int(np.sum((y_true == i) & (y_pred == j))) for j in (1, 2, 3)]
        for i in (1, 2, 3)
    ]
    per_class, f1_values = {}, []
    for category, label in zip((1, 2, 3), LABELS):
        tp = matrix[category - 1][category - 1]
        fp = sum(row[category - 1] for row in matrix) - tp
        fn = sum(matrix[category - 1]) - tp
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
        f1_values.append(f1)
        per_class[label] = {
            "precision": precision,
            "recall": recall,
            "f1": f1,
            "support": sum(matrix[category - 1]),
        }
    return {
        "accuracy": float(np.mean(y_true == y_pred)),
        "macro_f1": float(np.mean(f1_values)),
        "confusion_matrix_rows_true_cols_pred": matrix,
        "per_class": per_class,
    }


def leave_one_material_out(X, y, groups, model):
    folds, all_true, all_pred = {}, [], []
    for material in sorted(set(groups)):
        valid = groups == material
        pred, _ = fit_predict(model, X[~valid], y[~valid], X[valid])
        folds[material] = metrics(y[valid], pred)
        all_true.append(y[valid])
        all_pred.append(pred)
    return {
        "folds": folds,
        "aggregate": metrics(np.concatenate(all_true), np.concatenate(all_pred)),
    }


def main():
    X, y, groups, duplicate_removed = read_train()
    test_ids, X_test = read_test()
    train_idx, valid_idx = stratified_split(y)

    candidates = ["nearest_centroid", "gaussian_naive_bayes"]
    random_validation = {}
    material_validation = {}
    for model in candidates:
        valid_pred, _ = fit_predict(model, X[train_idx], y[train_idx], X[valid_idx])
        random_validation[model] = metrics(y[valid_idx], valid_pred)
        material_validation[model] = leave_one_material_out(X, y, groups, model)

    selected_model = max(
        candidates,
        key=lambda model: (
            material_validation[model]["aggregate"]["macro_f1"],
            random_validation[model]["macro_f1"],
        ),
    )
    if selected_model == "gaussian_naive_bayes":
        test_pred, test_distance = predict_gaussian_nb(
            X_test, fit_gaussian_nb(X, y)
        )
    else:
        test_pred, test_distance = predict_nearest_centroid(
            X_test, fit_nearest_centroid(X, y)
        )

    sorted_distance = np.sort(test_distance, axis=1)
    margins = sorted_distance[:, 1] - sorted_distance[:, 0]
    predictions = [
        {
            "序号": int(sample_id),
            "分类编号": int(category),
            "分类": LABELS[category - 1],
            "距离间隔": float(margins[i]),
        }
        for i, (sample_id, category) in enumerate(zip(test_ids, test_pred))
    ]
    with open(
        os.path.join(OUT_DIR, "predictions.csv"), "w", newline="", encoding="utf-8-sig"
    ) as file:
        writer = csv.DictWriter(file, fieldnames=predictions[0].keys())
        writer.writeheader()
        writer.writerows(predictions)

    counts = Counter(int(value) for value in test_pred)
    requested_ids = {1, 5, 15, 25, 35, 45, 55, 65, 75, 80}
    summary = {
        "method": "去均值/峰峰值归一化 + FFT/分位数/斜率特征",
        "selected_model": selected_model,
        "seed": SEED,
        "training_rows_after_duplicate_removal": int(len(y)),
        "duplicate_removed": duplicate_removed,
        "candidate_validation_metrics": random_validation,
        "leave_one_material_out_metrics": material_validation,
        "test_prediction_counts": {
            LABELS[c - 1]: int(counts.get(c, 0)) for c in (1, 2, 3)
        },
        "selected_predictions": [
            row for row in predictions if row["序号"] in requested_ids
        ],
    }
    with open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8") as file:
        json.dump(summary, file, ensure_ascii=False, indent=2)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
