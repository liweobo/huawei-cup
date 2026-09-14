import csv, json, math
from pathlib import Path
import openpyxl

BASE = Path(r'C:\Users\aaa\Desktop\text\pressure-test')
WORK = Path(r'C:\Users\aaa\Desktop\test')
PRED = WORK / 'outputs' / 'q1_model_comparison' / 'attachment2_model_predictions.csv'

def formula_or_number(v):
    return v

def numeric(v):
    return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(float(v))

def main():
    wb3 = openpyxl.load_workbook(BASE/'附件三（测试集）.xlsx', read_only=True, data_only=False)
    test3 = wb3.worksheets[0]
    wb3_values = openpyxl.load_workbook(BASE/'附件三（测试集）.xlsx', read_only=True, data_only=True)
    test3_values = wb3_values.worksheets[0]
    out4_formula = openpyxl.load_workbook(BASE/'附件四（Excel表）.xlsx', read_only=True, data_only=False).worksheets[0]
    out4_values = openpyxl.load_workbook(BASE/'附件四（Excel表）.xlsx', read_only=True, data_only=True).worksheets[0]
    with PRED.open(encoding='utf-8-sig', newline='') as f:
        preds=list(csv.DictReader(f))

    report={
        'attachment3': {'rows':test3.max_row-1,'columns':test3.max_column,'header_id':test3.cell(1,1).value,'first_ids_formula':[],'last_ids_formula':[],'first_ids_cached':[],'last_ids_cached':[],'first_meta':[],'last_meta':[]},
        'attachment4': {'rows':out4_formula.max_row-1,'columns':out4_formula.max_column,'headers':[out4_formula.cell(1,c).value for c in range(1,4)],'first_ids_formula':[],'last_ids_formula':[],'first_ids_cached':[],'last_ids_cached':[],'result_col2_nonempty_cached':0,'result_col3_nonempty_cached':0},
        'prediction_file': {'rows':len(preds),'first_ids':[],'last_ids':[],'first_values':[],'last_values':[]},
        'checks':{}
    }
    for r, row in enumerate(test3.iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        if r <= 6:
            report['attachment3']['first_ids_formula'].append(row[0]); report['attachment3']['first_meta'].append(list(row))
        elif r > test3.max_row - 5:
            report['attachment3']['last_ids_formula'].append(row[0]); report['attachment3']['last_meta'].append(list(row))
    for r, row in enumerate(test3_values.iter_rows(min_row=2, max_col=1, values_only=True), start=2):
        if r <= 6: report['attachment3']['first_ids_cached'].append(row[0])
        elif r > test3_values.max_row - 5: report['attachment3']['last_ids_cached'].append(row[0])
    for r in list(range(2,7))+list(range(out4_formula.max_row-4,out4_formula.max_row+1)):
        report['attachment4']['first_ids_formula' if r<7 else 'last_ids_formula'].append(out4_formula.cell(r,1).value)
        report['attachment4']['first_ids_cached' if r<7 else 'last_ids_cached'].append(out4_values.cell(r,1).value)
    for row in preds[:5]:
        report['prediction_file']['first_ids'].append(int(row['sample_id'])); report['prediction_file']['first_values'].append({k:row[k] for k in row if k.endswith('_id')})
    for row in preds[-5:]:
        report['prediction_file']['last_ids'].append(int(row['sample_id'])); report['prediction_file']['last_values'].append({k:row[k] for k in row if k.endswith('_id')})
    for r in range(2,out4_values.max_row+1):
        if out4_values.cell(r,2).value is not None: report['attachment4']['result_col2_nonempty_cached']+=1
        if out4_values.cell(r,3).value is not None: report['attachment4']['result_col3_nonempty_cached']+=1

    ids3=[]
    for row in test3_values.iter_rows(min_row=2, max_col=1, values_only=True): ids3.append(row[0])
    ids4_cached=[out4_values.cell(r,1).value for r in range(2,out4_values.max_row+1)]
    ids4_formula=[out4_formula.cell(r,1).value for r in range(2,out4_formula.max_row+1)]
    pred_ids=[int(x['sample_id']) for x in preds]
    report['checks']['attachment3_cached_ids_are_row_order_1_to_400'] = ids3 == list(range(1,401))
    report['checks']['attachment4_cached_ids_are_row_order_1_to_400'] = ids4_cached == list(range(1,401))
    report['checks']['attachment4_formula_ids_match_expected_pattern'] = ids4_formula[0] == 1 and all(isinstance(v,str) and v.startswith('=A') and '+1' in v for v in ids4_formula[1:])
    report['checks']['prediction_file_ids_are_1_to_80'] = pred_ids == list(range(1,81))
    report['checks']['q4_attachment3_and_attachment4_rows_match'] = (len(ids3)==len(ids4_cached)==400)
    report['checks']['q4_intended_write_mapping'] = 'prediction for attachment3 sample_id=i -> attachment4 row=i+1, column C'
    report['checks']['current_attachment4_is_unmodified_for_q4'] = report['attachment4']['result_col3_nonempty_cached'] == 0
    report['checks']['no_prediction_file_can_be_used_for_q4'] = 'The checked prediction file contains Q1 waveform labels for attachment2, not Q4 losses for attachment3.'
    # Explicitly flag the mismatch: Q1 prediction artifact is not a Q4 artifact.
    report['checks']['q4_loss_predictions_present'] = False
    report['checks']['q4_warning'] = 'Do not write q1_model_comparison/attachment2_model_predictions.csv into column C; it is the wrong dataset and wrong target.'
    (WORK/'q4_alignment_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
