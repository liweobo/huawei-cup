"""Question 1 baseline: simple waveform features + shallow decision tree.

Reads the original workbooks without modifying them and writes reproducible
metrics/predictions to outputs/q1_baseline/.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import joblib
import numpy as np
import openpyxl
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, f1_score
from sklearn.model_selection import StratifiedKFold, cross_val_score, train_test_split
from sklearn.tree import DecisionTreeClassifier, export_text


BASE = Path(r"C:\Users\aaa\Desktop\text\pressure-test")
TRAIN_FILE = BASE / "附件一（训练集）.xlsx"
TEST_FILE = BASE / "附件二（测试集）.xlsx"
OUT = Path(r"C:\Users\aaa\Desktop\test\outputs\q1_baseline")

LABEL_TO_ID = {"正弦波": 1, "三角波": 2, "梯形波": 3}
ID_TO_LABEL = {value: key for key, value in LABEL_TO_ID.items()}

FEATURE_NAMES = [
    "wave_std",
    "wave_rms",
    "extreme_ratio",
    "slope_cv",
    "slope_q90_q50",
    "flat_slope_ratio",
    "second_diff_ratio",
    "slope_max_mean_ratio",
]


def waveform_features(values) -> list[float]:
    """Extract amplitude-invariant shape descriptors from one 1024-point cycle."""
    x = np.asarray(values, dtype=float)
    if x.size != 1024 or not np.isfinite(x).all():
        raise ValueError(f"Expected 1024 finite waveform points, got {x.size}")

    x = x - x.mean()
    peak_to_peak = np.ptp(x)
    if peak_to_peak <= 0:
        raise ValueError("Constant waveform cannot be classified")
    x = x / peak_to_peak

    # Include the wrap-around difference because the samples represent a cycle.
    d = np.diff(np.r_[x, x[0]])
    abs_d = np.abs(d)
    eps = 1e-12
    q50, q90 = np.quantile(abs_d, [0.50, 0.90])
    max_d = abs_d.max()
    mean_d = abs_d.mean()
    dd = np.diff(np.r_[d, d[0]])

    return [
        float(x.std()),
        float(np.sqrt(np.mean(x * x))),
        float(np.mean(np.abs(x) >= 0.40)),
        float(abs_d.std() / (mean_d + eps)),
        float(q90 / (q50 + eps)),
        float(np.mean(abs_d <= 0.15 * max_d)),
        float(np.sqrt(np.mean(dd * dd)) / (np.sqrt(np.mean(d * d)) + eps)),
        float(max_d / (mean_d + eps)),
    ]


def load_training():
    workbook = openpyxl.load_workbook(TRAIN_FILE, read_only=True, data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        print(f"Reading training sheet: {sheet.title}", flush=True)
        for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = list(row)
            label = values[3]
            if label not in LABEL_TO_ID:
                raise ValueError(f"Unknown label {label!r} at {sheet.title}!{excel_row}")
            rows.append(
                {
                    "material": sheet.title,
                    "excel_row": excel_row,
                    "temperature": values[0],
                    "frequency": values[1],
                    "label_id": LABEL_TO_ID[label],
                    "label_name": label,
                    "features": waveform_features(values[4:1028]),
                }
            )
    workbook.close()
    return rows


def grouped_material_validation(rows):
    """Leave one material out, keeping the exact same shallow baseline."""
    results = []
    materials = sorted({row["material"] for row in rows})
    for held_out in materials:
        train = [row for row in rows if row["material"] != held_out]
        valid = [row for row in rows if row["material"] == held_out]
        x_train = np.asarray([row["features"] for row in train], dtype=float)
        y_train = np.asarray([row["label_id"] for row in train], dtype=int)
        x_valid = np.asarray([row["features"] for row in valid], dtype=float)
        y_valid = np.asarray([row["label_id"] for row in valid], dtype=int)
        grouped_model = DecisionTreeClassifier(
            max_depth=4,
            min_samples_leaf=10,
            class_weight="balanced",
            random_state=2024,
        )
        grouped_model.fit(x_train, y_train)
        pred = grouped_model.predict(x_valid)
        results.append(
            {
                "held_out_material": held_out,
                "samples": int(len(valid)),
                "accuracy": float(accuracy_score(y_valid, pred)),
                "macro_f1": float(f1_score(y_valid, pred, average="macro")),
                "confusion_matrix": confusion_matrix(y_valid, pred, labels=[1, 2, 3]).tolist(),
            }
        )
    return results


def load_test():
    workbook = openpyxl.load_workbook(TEST_FILE, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = []
    for sample_id, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        values = list(row)
        rows.append(
            {
                # Sequence cells after the first row contain formulas without cached
                # values in some readers, so use the required row order as the ID.
                "sample_id": sample_id,
                "temperature": values[1],
                "frequency": values[2],
                "material": values[3],
                "features": waveform_features(values[4:1028]),
            }
        )
    workbook.close()
    return rows


def write_feature_csv(path: Path, rows, include_label: bool):
    columns = ["sample_id" if not include_label else "material", "temperature", "frequency"]
    if include_label:
        columns += ["excel_row", "label_id", "label_name"]
    else:
        columns += ["material"]
    columns += FEATURE_NAMES
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            record = {key: row[key] for key in columns if key in row}
            record.update(dict(zip(FEATURE_NAMES, row["features"])))
            writer.writerow(record)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    train_rows = load_training()
    test_rows = load_test()

    X = np.asarray([row["features"] for row in train_rows], dtype=float)
    y = np.asarray([row["label_id"] for row in train_rows], dtype=int)

    X_train, X_valid, y_train, y_valid = train_test_split(
        X, y, test_size=0.20, random_state=2024, stratify=y
    )
    model = DecisionTreeClassifier(
        max_depth=4,
        min_samples_leaf=10,
        class_weight="balanced",
        random_state=2024,
    )
    model.fit(X_train, y_train)
    valid_pred = model.predict(X_valid)

    cv = StratifiedKFold(n_splits=5, shuffle=True, random_state=2024)
    cv_accuracy = cross_val_score(model, X, y, cv=cv, scoring="accuracy")
    cv_macro_f1 = cross_val_score(model, X, y, cv=cv, scoring="f1_macro")

    report = classification_report(
        y_valid,
        valid_pred,
        labels=[1, 2, 3],
        target_names=[ID_TO_LABEL[i] for i in [1, 2, 3]],
        output_dict=True,
        digits=6,
        zero_division=0,
    )
    matrix = confusion_matrix(y_valid, valid_pred, labels=[1, 2, 3])

    # Refit the same fixed baseline on all labeled data for final test predictions.
    model.fit(X, y)
    X_test = np.asarray([row["features"] for row in test_rows], dtype=float)
    test_pred = model.predict(X_test)
    test_prob = model.predict_proba(X_test)

    counts = {ID_TO_LABEL[i]: int(np.sum(test_pred == i)) for i in [1, 2, 3]}
    selected_ids = [1, 5, 15, 25, 35, 45, 55, 65, 75, 80]
    selected = []
    for sample_id in selected_ids:
        pred = int(test_pred[sample_id - 1])
        selected.append(
            {"sample_id": sample_id, "prediction_id": pred, "prediction_name": ID_TO_LABEL[pred]}
        )

    metrics = {
        "model": "DecisionTreeClassifier(max_depth=4, min_samples_leaf=10, class_weight='balanced')",
        "features": FEATURE_NAMES,
        "train_samples": int(len(y)),
        "holdout_samples": int(len(y_valid)),
        "holdout_accuracy": float(accuracy_score(y_valid, valid_pred)),
        "holdout_macro_f1": float(f1_score(y_valid, valid_pred, average="macro")),
        "holdout_confusion_matrix_labels": ["正弦波", "三角波", "梯形波"],
        "holdout_confusion_matrix": matrix.tolist(),
        "holdout_classification_report": report,
        "cv_accuracy_scores": cv_accuracy.tolist(),
        "cv_accuracy_mean": float(cv_accuracy.mean()),
        "cv_accuracy_std": float(cv_accuracy.std()),
        "cv_macro_f1_scores": cv_macro_f1.tolist(),
        "cv_macro_f1_mean": float(cv_macro_f1.mean()),
        "cv_macro_f1_std": float(cv_macro_f1.std()),
        "leave_one_material_out": grouped_material_validation(train_rows),
        "test_samples": int(len(test_rows)),
        "test_prediction_counts": counts,
        "selected_test_predictions": selected,
    }

    with (OUT / "metrics.json").open("w", encoding="utf-8") as handle:
        json.dump(metrics, handle, ensure_ascii=False, indent=2)

    with (OUT / "attachment2_predictions.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        fields = ["sample_id", "prediction_id", "prediction_name", "confidence"]
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for i, pred in enumerate(test_pred):
            writer.writerow(
                {
                    "sample_id": i + 1,
                    "prediction_id": int(pred),
                    "prediction_name": ID_TO_LABEL[int(pred)],
                    "confidence": round(float(test_prob[i].max()), 6),
                }
            )

    write_feature_csv(OUT / "train_features.csv", train_rows, include_label=True)
    write_feature_csv(OUT / "test_features.csv", test_rows, include_label=False)
    (OUT / "tree_rules.txt").write_text(
        export_text(model, feature_names=FEATURE_NAMES, decimals=6), encoding="utf-8"
    )
    joblib.dump(model, OUT / "decision_tree.joblib")

    summary = [
        "# Q1 waveform classification baseline",
        "",
        f"- Model: {metrics['model']}",
        f"- Training samples: {len(y)}",
        f"- Holdout accuracy: {metrics['holdout_accuracy']:.6f}",
        f"- Holdout Macro-F1: {metrics['holdout_macro_f1']:.6f}",
        f"- 5-fold CV accuracy: {metrics['cv_accuracy_mean']:.6f} +/- {metrics['cv_accuracy_std']:.6f}",
        f"- 5-fold CV Macro-F1: {metrics['cv_macro_f1_mean']:.6f} +/- {metrics['cv_macro_f1_std']:.6f}",
        f"- Test counts: {counts}",
        "",
        "## Holdout confusion matrix",
        "",
        "Rows are true labels and columns are predicted labels: sine, triangle, trapezoid.",
        "",
        "```",
        *[str(row) for row in matrix.tolist()],
        "```",
        "",
        "## Selected test samples",
        "",
        "| sample_id | prediction_id | prediction_name |",
        "|---:|---:|---|",
        *[f"| {r['sample_id']} | {r['prediction_id']} | {r['prediction_name']} |" for r in selected],
    ]
    (OUT / "summary.md").write_text("\n".join(summary) + "\n", encoding="utf-8")

    print(json.dumps(metrics, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
