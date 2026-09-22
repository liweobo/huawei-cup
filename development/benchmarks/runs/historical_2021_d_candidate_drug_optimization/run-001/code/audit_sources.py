"""Read-only source, workbook, and entity audit for the 2021D blind run."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


RUN_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = RUN_ROOT / "source-provenance" / "original"
OUTPUT_DIR = RUN_ROOT / "outputs" / "source-audit"

FILES = [
    "抗胰腺癌候选药物的优化建模.docx",
    "ERα_activity.xlsx",
    "ADMET.xlsx",
    "Molecular_Descriptor.xlsx",
    "分子描述符含义解释.xlsx",
]
MISSING_STRINGS = {"", "na", "n/a", "nan", "null", "none", "missing", "-"}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def blank_or_marker(value: Any) -> bool:
    if value is None:
        return True
    return isinstance(value, str) and value.strip().lower() in MISSING_STRINGS


def workbook_audit(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=False)
    workbook_record: dict[str, Any] = {
        "filename": path.name,
        "sha256": sha256(path),
        "size_bytes": path.stat().st_size,
        "sheet_names": wb.sheetnames,
        "hidden_sheets": [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"],
        "sheets": [],
    }
    for ws in wb.worksheets:
        headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        formula_cells = []
        missing_cells = 0
        explicit_missing_markers: Counter[str] = Counter()
        duplicate_full_rows = 0
        row_signatures: Counter[tuple[Any, ...]] = Counter()
        ids = []
        for r, cells in enumerate(ws.iter_rows(min_row=2), 2):
            row = tuple(cell.value for cell in cells)
            row_signatures[row] += 1
            ids.append(row[0])
            for cell, value in zip(cells, row):
                if blank_or_marker(value):
                    missing_cells += 1
                    if isinstance(value, str):
                        explicit_missing_markers[value.strip()] += 1
                if cell.data_type == "f":
                    formula_cells.append(cell.coordinate)
        duplicate_full_rows = sum(count - 1 for count in row_signatures.values() if count > 1)
        is_entity_sheet = bool(headers and headers[0] == "SMILES")
        id_counts = Counter(ids) if is_entity_sheet else Counter()
        duplicate_ids = sorted(str(key) for key, count in id_counts.items() if count > 1)
        workbook_record["sheets"].append(
            {
                "name": ws.title,
                "state": ws.sheet_state,
                "dimensions": {"rows_including_header": ws.max_row, "columns": ws.max_column},
                "header_row": 1,
                "headers": headers,
                "formula_cell_count": len(formula_cells),
                "formula_cell_examples": formula_cells[:20],
                "missing_cells_excluding_header": missing_cells,
                "explicit_missing_markers": dict(explicit_missing_markers),
                "id_column": headers[0] if is_entity_sheet else None,
                "duplicate_id_count": len(duplicate_ids) if is_entity_sheet else None,
                "duplicate_ids": duplicate_ids[:50] if is_entity_sheet else [],
                "duplicate_full_rows": duplicate_full_rows,
            }
        )
    return workbook_record


def load_ids(path: Path, sheet: str) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    return [str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True)]


def entity_alignment() -> dict[str, Any]:
    sources = {
        "activity": SOURCE_DIR / "ERα_activity.xlsx",
        "admet": SOURCE_DIR / "ADMET.xlsx",
        "descriptor": SOURCE_DIR / "Molecular_Descriptor.xlsx",
    }
    result: dict[str, Any] = {"key": "SMILES", "splits": {}}
    for split in ("training", "test"):
        ids = {name: load_ids(path, split) for name, path in sources.items()}
        sets = {name: set(values) for name, values in ids.items()}
        reference = ids["activity"]
        result["splits"][split] = {
            "row_counts": {name: len(values) for name, values in ids.items()},
            "unique_counts": {name: len(set(values)) for name, values in ids.items()},
            "duplicate_counts": {
                name: sum(count - 1 for count in Counter(values).values() if count > 1)
                for name, values in ids.items()
            },
            "set_equality": all(values == sets["activity"] for values in sets.values()),
            "row_order_equality": all(values == reference for values in ids.values()),
            "missing_from_activity": {
                name: sorted(values - sets["activity"])[:50] for name, values in sets.items()
            },
            "missing_from_source": {
                name: sorted(sets["activity"] - values)[:50] for name, values in sets.items()
            },
        }
    training_ids = set(load_ids(sources["activity"], "training"))
    test_ids = set(load_ids(sources["activity"], "test"))
    result["train_test_overlap_count"] = len(training_ids & test_ids)
    result["train_test_overlap_examples"] = sorted(training_ids & test_ids)[:50]
    result["status"] = (
        "PASS"
        if all(
            result["splits"][split]["set_equality"]
            and result["splits"][split]["row_order_equality"]
            and all(v == 0 for v in result["splits"][split]["duplicate_counts"].values())
            for split in ("training", "test")
        )
        and result["train_test_overlap_count"] == 0
        else "FAIL"
    )
    return result


def activity_transform_check() -> dict[str, Any]:
    wb = load_workbook(SOURCE_DIR / "ERα_activity.xlsx", read_only=True, data_only=True)
    ws = wb["training"]
    deviations = []
    invalid = []
    for row_number, (_, ic50, pic50) in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if ic50 is None or pic50 is None or float(ic50) <= 0:
            invalid.append(row_number)
            continue
        expected = 9.0 - math.log10(float(ic50))
        deviations.append(abs(float(pic50) - expected))
    return {
        "definition_checked": "pIC50 = 9 - log10(IC50_nM)",
        "rows_checked": len(deviations),
        "invalid_rows": invalid,
        "max_absolute_deviation": max(deviations) if deviations else None,
        "status": "PASS" if deviations and max(deviations) < 1e-10 and not invalid else "FAIL",
    }


def descriptor_dictionary() -> dict[str, Any]:
    wb = load_workbook(SOURCE_DIR / "分子描述符含义解释.xlsx", read_only=True, data_only=True)
    ws = wb["Detailed"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    mapping = {}
    for java_class, descriptor, description, klass in rows:
        if descriptor:
            mapping[str(descriptor)] = {
                "java_class": java_class,
                "description": description,
                "class": klass,
            }
    descriptor_wb = load_workbook(
        SOURCE_DIR / "Molecular_Descriptor.xlsx", read_only=True, data_only=True
    )
    headers = list(next(descriptor_wb["training"].iter_rows(min_row=1, max_row=1, values_only=True)))[1:]
    casefold_mapping = {name.casefold(): (name, value) for name, value in mapping.items()}
    uncovered = [str(name) for name in headers if str(name).casefold() not in casefold_mapping]
    extra = [name for name in mapping if name not in set(map(str, headers))]
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with (OUTPUT_DIR / "descriptor-dictionary.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["descriptor", "java_class", "description", "class"])
        writer.writeheader()
        for descriptor in map(str, headers):
            item = casefold_mapping.get(descriptor.casefold(), (None, {}))[1]
            writer.writerow({"descriptor": descriptor, **item})
    return {
        "descriptor_columns": len(headers),
        "detailed_dictionary_entries": len(mapping),
        "covered_descriptor_columns": len(headers) - len(uncovered),
        "uncovered_descriptor_columns": uncovered,
        "dictionary_only_entries": extra,
        "status": "PASS" if not uncovered else "PARTIAL",
    }


def extract_problem() -> dict[str, Any]:
    path = SOURCE_DIR / "抗胰腺癌候选药物的优化建模.docx"
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs) + "\n"
    (RUN_ROOT / "source-provenance" / "problem-text.txt").write_text(text, encoding="utf-8")
    return {
        "filename": path.name,
        "core_title_property": doc.core_properties.title or None,
        "first_paragraph": paragraphs[0] if paragraphs else None,
        "document_title": paragraphs[1] if len(paragraphs) > 1 else None,
        "paragraph_count_nonempty": len(paragraphs),
        "table_count": len(doc.tables),
        "contains_breast_cancer": "乳腺癌" in text,
        "contains_pancreatic_cancer": "胰腺癌" in text,
        "contains_er_alpha": "ERα" in text,
        "q1_q4_present": all(
            f"问题 {number}" in text or f"问题{number}" in text for number in range(1, 5)
        ),
    }


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((RUN_ROOT / "source-provenance" / "source-manifest.json").read_text(encoding="utf-8"))
    expected = {item["filename"]: item for item in manifest["files"]}
    file_checks = []
    for filename in FILES:
        path = SOURCE_DIR / filename
        record = expected[filename]
        file_checks.append(
            {
                "filename": filename,
                "exists": path.exists(),
                "size_bytes": path.stat().st_size,
                "size_matches_manifest": path.stat().st_size == record["size_bytes"],
                "sha256": sha256(path),
                "sha256_matches_manifest": sha256(path) == record["sha256"],
                "git_blob_verified_at_acquisition": record["verification"],
            }
        )
    audit = {
        "file_checks": file_checks,
        "problem_document": extract_problem(),
        "workbooks": [workbook_audit(SOURCE_DIR / name) for name in FILES if name.endswith(".xlsx")],
        "activity_transform": activity_transform_check(),
        "descriptor_dictionary": descriptor_dictionary(),
        "entity_alignment": entity_alignment(),
    }
    (OUTPUT_DIR / "source-audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (RUN_ROOT / "source-provenance" / "workbook-integrity.json").write_text(
        json.dumps(audit["workbooks"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (RUN_ROOT / "outputs" / "entity-alignment.json").write_text(
        json.dumps(audit["entity_alignment"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
